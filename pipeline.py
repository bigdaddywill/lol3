from __future__ import annotations

import argparse
import json
import re
from dataclasses import asdict, dataclass, field
from html.parser import HTMLParser
from pathlib import Path
from typing import Any
from urllib.parse import quote_plus, urljoin, urlparse
from urllib.parse import parse_qs
import base64
from urllib.request import Request, urlopen
from zipfile import ZipFile
from xml.etree import ElementTree as ET

NS = {
    "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "rel": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}

def _col_index(ref: str) -> int:
    m = re.match(r"([A-Z]+)", ref.upper())
    if not m:
        return 0
    n = 0
    for ch in m.group(1):
        n = n * 26 + ord(ch) - 64
    return n - 1

def _text(node: ET.Element | None) -> str:
    return "" if node is None else "".join(node.itertext())

@dataclass
class Sheet:
    name: str
    rows: list[list[str]]

def read_xlsx(path: str | Path) -> list[Sheet]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to read XLSX workbooks") from exc

    wb = load_workbook(filename=path, read_only=True, data_only=True)
    sheets: list[Sheet] = []
    try:
        for ws in wb.worksheets:
            rows: list[list[str]] = []
            for row in ws.iter_rows(values_only=True):
                rows.append([norm(v) for v in row])
            sheets.append(Sheet(ws.title, rows))
    finally:
        wb.close()
    return sheets

TARGET_SHEET = "on leads magnet"

ALIASES = {
    "company": ["company", "company name", "business", "business name", "contractor", "contractor name", "legal name"],
    "contact": ["contact", "contact name", "primary contact", "owner", "poc", "point of contact", "contact person"],
    "email": ["email", "email address", "e-mail", "email 1", "email1"],
    "phone": ["phone", "phone number", "telephone", "mobile", "cell", "phone 1", "phone1"],
    "website": ["website", "web", "url", "site", "company website"],
    "city": ["city", "town"],
    "state": ["state", "state code", "province", "st"],
    "zip": ["zip", "zip code", "postal", "postal code"],
    "trade": ["trade", "specialty", "specialties", "services", "service", "industry", "category", "work type"],
}

def norm(value: Any) -> str:
    s = "" if value is None else str(value).strip()
    return "" if s.lower() in {"nan", "none", "nat"} else s

def key(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", " ", norm(value).lower()).strip()

def pick(headers: list[str], aliases: list[str]) -> str | None:
    exact = {key(h): h for h in headers}
    for alias in aliases:
        if key(alias) in exact:
            return exact[key(alias)]
    for header in headers:
        h = key(header)
        if any(key(alias) in h for alias in aliases):
            return header
    return None

@dataclass
class Contractor:
    sheet: str
    company: str
    contact: str = ""
    email: str = ""
    phone: str = ""
    website: str = ""
    city: str = ""
    state: str = ""
    zip: str = ""
    trade: str = ""
    def to_dict(self) -> dict[str, str]:
        return asdict(self)

def load_contractors(path: str | Path) -> tuple[list[Contractor], dict[str, Any]]:
    sheets = read_xlsx(path)
    contractors: list[Contractor] = []
    seen: set[tuple[str, str, str]] = set()
    audit: dict[str, Any] = {
        "sheets": [{"name": s.name, "rows": len(s.rows)} for s in sheets],
        "rows_seen": 0,
        "rows_loaded": 0,
        "blank_rows": 0,
        "duplicates_removed": 0,
        "header_mapping": {},
    }
    for sheet in sheets:
        if key(sheet.name) != TARGET_SHEET:
            continue
        if not sheet.rows:
            audit["header_mapping"][sheet.name] = {}
            continue
        headers = [norm(x) for x in sheet.rows[0]]
        mapping = {f: pick(headers, aliases) for f, aliases in ALIASES.items()}
        audit["header_mapping"][sheet.name] = mapping
        audit["rows_seen"] += max(0, len(sheet.rows) - 1)
        for row in sheet.rows[1:]:
            values = [norm(x) for x in row]
            if not any(values):
                audit["blank_rows"] += 1
                continue
            raw = {h: values[i] if i < len(values) else "" for i, h in enumerate(headers) if h}
            def get(field_name: str) -> str:
                col = mapping[field_name]
                return raw.get(col, "") if col else ""
            company = get("company") or next((v for v in values if v), "")
            if not company:
                continue
            c = Contractor(
                sheet=sheet.name,
                company=company,
                contact=get("contact"),
                email=get("email"),
                phone=get("phone"),
                website=get("website"),
                city=get("city"),
                state=get("state"),
                zip=get("zip"),
                trade=get("trade"),
            )
            fingerprint = (c.company.lower(), c.email.lower(), re.sub(r"\D", "", c.phone))
            if fingerprint in seen:
                audit["duplicates_removed"] += 1
                continue
            seen.add(fingerprint)
            contractors.append(c)
    audit["rows_loaded"] = len(contractors)
    audit["rows_with_email"] = sum(bool(c.email) for c in contractors)
    audit["rows_with_phone"] = sum(bool(c.phone) for c in contractors)
    audit["rows_with_location"] = sum(bool(c.city and c.state) for c in contractors)
    audit["rows_with_trade"] = sum(bool(c.trade) for c in contractors)
    return contractors, audit

class HTMLTextParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self.title: list[str] = []
        self.in_title = False
        self.skip_depth = 0
    def handle_starttag(self, tag: str, attrs):
        if tag in {"script", "style", "noscript", "svg"}:
            self.skip_depth += 1
        if tag == "title":
            self.in_title = True
    def handle_endtag(self, tag: str):
        if tag == "title":
            self.in_title = False
        if tag in {"script", "style", "noscript", "svg"} and self.skip_depth:
            self.skip_depth -= 1
    def handle_data(self, data: str):
        if self.skip_depth:
            return
        text = re.sub(r"\s+", " ", data).strip()
        if text:
            self.parts.append(text)
            if self.in_title:
                self.title.append(text)

def fetch_page(url: str, timeout: int = 20) -> tuple[str, str]:
    req = Request(url, headers={"User-Agent": "CommercialBidPipeline/1.0"})
    with urlopen(req, timeout=timeout) as resp:
        raw = resp.read(3_000_000)
        charset = resp.headers.get_content_charset() or "utf-8"
    html = raw.decode(charset, errors="replace")
    parser = HTMLTextParser()
    parser.feed(html)
    return " ".join(parser.title).strip(), " ".join(parser.parts)

def unwrap_bing_url(url: str) -> str:
    try:
        parsed = urlparse(url)
        if "bing.com" not in parsed.netloc:
            return url
        raw = parse_qs(parsed.query).get("u", [""])[0]
        if not raw:
            return url
        token = raw[2:] if raw.startswith("a1") else raw
        decoded = base64.urlsafe_b64decode(token + "=" * (-len(token) % 4)).decode("utf-8", "ignore")
        return decoded if decoded.startswith(("http://", "https://")) else url
    except Exception:
        return url

@dataclass
class SearchResult:
    title: str
    url: str
    snippet: str = ""

def _extract_anchor_results(html: str) -> list[SearchResult]:
    results: list[SearchResult] = []
    pattern = re.compile(
        r"""<a[^>]*href=["']([^"']+)["'][^>]*>(.*?)</a>""",
        flags=re.I | re.S,
    )
    for href, title_html in pattern.findall(html):
        if "result__a" not in title_html.lower() and not re.search(r"<h2", title_html, re.I):
            continue
        title = re.sub(r"\\s+", " ", re.sub(r"<[^>]+>", " ", title_html)).strip()
        if not title or len(title) < 4:
            continue
        results.append(SearchResult(title=title, url=href))
    return results

class SearchHTMLParser(HTMLParser):
    def __init__(self, mode: str):
        super().__init__()
        self.mode = mode
        self.in_h2 = False
        self.in_anchor = False
        self.anchor_href = ""
        self.title_parts: list[str] = []
        self.results: list[SearchResult] = []
        self.anchor_is_result = False

    def handle_starttag(self, tag: str, attrs):
        attrs_dict = dict(attrs)
        if tag.lower() == "h2":
            self.in_h2 = True
        if tag.lower() == "a":
            classes = set((attrs_dict.get("class") or "").split())
            self.in_anchor = True
            self.anchor_href = attrs_dict.get("href") or ""
            self.title_parts = []
            if self.mode == "bing":
                self.anchor_is_result = self.in_h2
            else:
                self.anchor_is_result = "result__a" in classes

    def handle_data(self, data: str):
        if self.in_anchor and self.anchor_is_result:
            self.title_parts.append(data)

    def handle_endtag(self, tag: str):
        tag = tag.lower()
        if tag == "a" and self.in_anchor:
            if self.anchor_is_result and self.anchor_href:
                title = re.sub(r"\\s+", " ", " ".join(self.title_parts)).strip()
                if title:
                    self.results.append(SearchResult(title=title, url=unwrap_bing_url(self.anchor_href)))
            self.in_anchor = False
            self.anchor_href = ""
            self.title_parts = []
            self.anchor_is_result = False
        elif tag == "h2":
            self.in_h2 = False


def search_bing(query: str, max_results: int = 10, timeout: int = 20) -> list[SearchResult]:
    search_url = "https://www.bing.com/search?q=" + quote_plus(query)
    req = Request(search_url, headers={"User-Agent": "Mozilla/5.0 (commercial bid research)"})
    with urlopen(req, timeout=timeout) as resp:
        html = resp.read(3_000_000).decode(resp.headers.get_content_charset() or "utf-8", errors="replace")
    parser = SearchHTMLParser("bing")
    parser.feed(html)
    return parser.results[:max_results]

def search_duckduckgo(query: str, max_results: int = 10, timeout: int = 20) -> list[SearchResult]:
    search_url = "https://html.duckduckgo.com/html/?q=" + quote_plus(query)
    req = Request(search_url, headers={"User-Agent": "Mozilla/5.0 (commercial bid research)"})
    with urlopen(req, timeout=timeout) as resp:
        html = resp.read(3_000_000).decode(resp.headers.get_content_charset() or "utf-8", errors="replace")
    parser = SearchHTMLParser("ddg")
    parser.feed(html)
    return parser.results[:max_results]

def search_web(query: str, max_results: int = 10) -> list[SearchResult]:
    seen: set[str] = set()
    out: list[SearchResult] = []
    for fn in (search_bing, search_duckduckgo):
        try:
            for result in fn(query, max_results=max_results):
                if result.url in seen:
                    continue
                seen.add(result.url)
                out.append(result)
                if len(out) >= max_results:
                    return out
        except Exception:
            continue
    return out

@dataclass
class SourceField:
    value: str
    source_url: str
    confidence: str = "high"

@dataclass
class Bid:
    title: str
    source_url: str
    source_domain: str
    bid_number: SourceField | None = None
    scope: SourceField | None = None
    location: SourceField | None = None
    due_date: SourceField | None = None
    contact_name: SourceField | None = None
    contact_email: SourceField | None = None
    contact_phone: SourceField | None = None
    raw_excerpt: str = ""
    tags: list[str] = field(default_factory=list)
    def field(self, name: str) -> str:
        value = getattr(self, name, None)
        return value.value if isinstance(value, SourceField) else ""
    def to_dict(self) -> dict[str, Any]:
        return asdict(self)

def source_field(value: str, url: str, confidence: str = "high") -> SourceField | None:
    value = (value or "").strip()
    return SourceField(value, url, confidence) if value else None

def find(patterns: list[str], text: str) -> str:
    for pattern in patterns:
        match = re.search(pattern, text, flags=re.I)
        if match:
            return match.group(1).strip(" .,:;-|")
    return ""

def extract_bid(result: SearchResult, page_title: str, page_text: str) -> Bid:
    combined = (result.snippet + " " + page_text).strip()
    due = find([
        r"(?:bid|proposal|solicitation|submission)[^.;]{0,100}?(?:due|deadline)[^.;]{0,60}?((?:Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?|Jul(?:y)?|Aug(?:ust)?|Sep(?:t(?:ember)?)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?)\s+\d{1,2},?\s+\d{4}(?:\s+at\s+[0-9: AMP]+)?)",
        r"(?:due date|bid due|deadline)\s*[:\-]?\s*(\d{1,2}/\d{1,2}/\d{2,4})",
    ], combined)
    bid_no = find([r"(?:bid|solicitation|project|opportunity|rfp|ifb)\s*(?:number|no\.|#)?\s*[:#-]?\s*([A-Z0-9][A-Z0-9._/-]{2,})"], combined)
    email = find([r"([A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,})"], combined)
    phone = find([r"(\+?1?[\s.-]?\(?\d{3}\)?[\s.-]?\d{3}[\s.-]?\d{4})"], combined)
    contact = find([r"(?:contact|point of contact|poc)\s*[:\-]\s*([A-Z][A-Za-z.' -]{2,60})"], combined)
    location = find([r"(?:project location|job site|site address|location)\s*[:\-]\s*([^.;|]{4,140})"], combined)
    scope = find([r"(?:scope of work|scope|project description|work includes)\s*[:\-]\s*([^.;|]{30,900})"], combined)
    return Bid(
        title=page_title or result.title,
        source_url=result.url,
        source_domain=urlparse(result.url).netloc,
        bid_number=source_field(bid_no, result.url),
        scope=source_field(scope or result.snippet, result.url, "medium"),
        location=source_field(location, result.url),
        due_date=source_field(due, result.url),
        contact_name=source_field(contact, result.url),
        contact_email=source_field(email.lower(), result.url),
        contact_phone=source_field(phone, result.url),
        raw_excerpt=combined[:2000],
    )

def discover(queries: list[str] | str, max_results: int = 10) -> list[Bid]:
    if isinstance(queries, str):
        queries = [queries]
    results: list[SearchResult] = []
    seen: set[str] = set()
    per_query = max(3, max_results // max(1, len(queries)))
    for query in queries:
        for result in search_web(query, max_results=per_query):
            if result.url in seen:
                continue
            seen.add(result.url)
            results.append(result)
            if len(results) >= max_results:
                break
        if len(results) >= max_results:
            break

    bids: list[Bid] = []
    for result in results:
        try:
            title, text = fetch_page(result.url)
            bids.append(extract_bid(result, title, text))
        except Exception:
            bids.append(extract_bid(result, result.title, result.snippet))
    return bids

TRADE_GROUPS = {
    "landscaping": {"landscape", "landscaping", "grounds", "tree", "arbor", "mowing", "mulch", "pruning"},
    "electrical": {"electrical", "electric", "lighting", "low voltage"},
    "plumbing": {"plumbing", "plumber", "pipe", "piping", "drain"},
    "hvac": {"hvac", "mechanical", "air conditioning", "heating", "cooling"},
    "roofing": {"roof", "roofing", "reroof"},
    "painting": {"paint", "painting", "coating"},
    "flooring": {"floor", "flooring", "tile", "carpet", "vinyl"},
    "general": {"general contractor", "gc", "remodel", "renovation", "construction"},
}

def tokens(text: str) -> set[str]:
    return {x for x in re.findall(r"[a-z0-9]+", (text or "").lower()) if len(x) >= 3}

def trade_terms(trade: str) -> set[str]:
    base = tokens(trade)
    terms = set(base)
    for group in TRADE_GROUPS.values():
        if base & group:
            terms |= group
    return terms

@dataclass
class Match:
    contractor: Contractor
    bid: Bid
    score: int
    reasons: list[str]
    def to_dict(self) -> dict[str, Any]:
        return asdict(self)

def score_match(contractor: Contractor, bid: Bid) -> Match:
    hay = " ".join([bid.title, bid.field("scope"), bid.field("location"), bid.raw_excerpt]).lower()
    if not any(term in hay for term in ("invitation to bid", "request for bid", "request for proposal", "rfp", "rfq", "ifb", "solicitation", "bid opportunity", "procurement")):
        return Match(contractor, bid, 0, ["not clearly a bid opportunity"])
    score = 0
    reasons: list[str] = []
    if contractor.state and re.search(rf"\b{re.escape(contractor.state.lower())}\b", hay):
        score += 30
        reasons.append("state match")
    if contractor.city and re.search(rf"\b{re.escape(contractor.city.lower())}\b", hay):
        score += 25
        reasons.append("city match")
    hits = sorted(term for term in trade_terms(contractor.trade) if term in hay)
    if hits:
        score += min(45, len(hits) * 10)
        reasons.append("trade match: " + ", ".join(hits[:5]))
    if bid.field("due_date"):
        score += 5
        reasons.append("deadline present")
    return Match(contractor, bid, score, reasons)

def rank_matches(contractor: Contractor, bids: list[Bid], minimum_score: int = 30) -> list[Match]:
    matches = [score_match(contractor, bid) for bid in bids]
    return sorted([m for m in matches if m.score >= minimum_score], key=lambda m: (-m.score, m.bid.title.lower()))

def clean(text: str) -> str:
    return " ".join((text or "").split())

def generate_email(primary: Match, related: list[Match] | None = None) -> tuple[str, str]:
    c = primary.contractor
    matches = [primary] + list(related or [])
    subject = f"Open Commercial Remodeling Bids in {c.state or 'Your Area'}"
    greeting = c.contact or c.company or "there"
    lines = [f"Hi {greeting},", "", f"William here — we found {len(matches)} active commercial bid opportunities that may fit {c.company or 'your company'}.", ""]
    for i, m in enumerate(matches, 1):
        b = m.bid
        lines += [
            f"{i}. {b.title}",
            f"Scope: {clean(b.field('scope')) or 'Not stated in the source'}",
            f"Location: {clean(b.field('location')) or 'Not stated in the source'}",
            f"Bid due: {clean(b.field('due_date')) or 'Not stated in the source'}",
            f"POC: {clean(b.field('contact_name')) or 'Not stated in the source'}",
            f"POC email: {clean(b.field('contact_email')) or 'Not stated in the source'}",
            f"POC phone: {clean(b.field('contact_phone')) or 'Not stated in the source'}",
            f"Source: {b.source_url}",
            "",
        ]
    lines += ["We have attached an invitation to bid to the opportunity we think is the closest fit.", "If you'd like help filling it out or have questions, call us at 855-799-8420.", "", "Best,", "William"]
    return subject, "\n".join(lines)

def generate_sms(primary: Match) -> str:
    c = primary.contractor
    b = primary.bid
    text = f"Hi {c.contact or c.company or 'there'}, William here — I just sent you an active commercial bid we think may be a fit. {b.title}."
    if b.field("location"):
        text += f" Location: {clean(b.field('location'))}."
    if b.field("due_date"):
        text += f" Bid due: {clean(b.field('due_date'))}."
    if c.email:
        text += f" Sent to {c.email}."
    text += " Take a look when you have a chance. Questions: 855-799-8420."
    return text[:480]

def write_artifacts(primary: Match, out_dir: str | Path, related: list[Match] | None = None) -> Path:
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    subject, email = generate_email(primary, related)
    sms = generate_sms(primary)
    payload = {
        "contractor": primary.contractor.to_dict(),
        "match": {"score": primary.score, "reasons": primary.reasons},
        "bid": primary.bid.to_dict(),
        "verification": {"source_url": primary.bid.source_url, "human_review_before_send": True},
        "email": {"subject": subject, "body": email},
        "sms": sms,
    }
    (out / "bid_packet.json").write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    invitation = [
        "# Invitation to Bid", "", f"## {primary.bid.title}", "",
        f"**Contractor:** {primary.contractor.company}", f"**Match score:** {primary.score}",
        "", "### Scope", clean(primary.bid.field("scope")) or "Not stated in source.",
        "", "### Location", clean(primary.bid.field("location")) or "Not stated in source.",
        "", "### Bid Due", clean(primary.bid.field("due_date")) or "Not stated in source.",
        "", "### Point of Contact",
        f"Name: {primary.bid.field('contact_name') or 'Not stated in source'}",
        f"Email: {primary.bid.field('contact_email') or 'Not stated in source'}",
        f"Phone: {primary.bid.field('contact_phone') or 'Not stated in source'}",
        "", "### Source", primary.bid.source_url, "", "_Human verification required before outreach._"
    ]
    (out / "invitation.md").write_text("\n".join(invitation), encoding="utf-8")
    (out / "email.txt").write_text(f"Subject: {subject}\n\n{email}", encoding="utf-8")
    (out / "sms.txt").write_text(sms, encoding="utf-8")
    return out

def inspect_workbook(path: str) -> None:
    contractors, audit = load_contractors(path)
    print(json.dumps({"audit": audit, "sample": [c.to_dict() for c in contractors[:10]]}, indent=2, ensure_ascii=False))

def run(args: argparse.Namespace) -> None:
    contractors, audit = load_contractors(args.input)
    if args.state:
        contractors = [c for c in contractors if c.state.lower().strip() == args.state.lower().strip()]
    contractors = contractors[:args.limit]
    output = Path(args.output)
    output.mkdir(parents=True, exist_ok=True)
    manifest: list[dict[str, Any]] = []
    for index, contractor in enumerate(contractors, 1):
        trade = contractor.trade
        state = contractor.state
        city = contractor.city
        base = '"commercial bid" OR "invitation to bid" OR "request for bid"'
        queries = [
            " ".join(x for x in [base, trade, city, state, "commercial construction"] if x),
            " ".join(x for x in ['"invitation to bid"', trade, state] if x),
            " ".join(x for x in ['"request for proposal"', trade, state] if x),
            " ".join(x for x in ['site:gov', trade, state, bid] if x) if False else " ".join(x for x in ['site:gov', trade, state, "bid"] if x),
        ]
        if args.query:
            queries = [args.query]
        try:
            bids = discover(queries, max_results=args.max_results)
        except Exception as exc:
            manifest.append({"company": contractor.company, "status": "discovery_error", "error": str(exc), "query": query})
            continue
        matches = rank_matches(contractor, bids, minimum_score=args.minimum_score)
        if not matches:
            manifest.append({
                "company": contractor.company,
                "status": "no_match",
                "queries": queries,
                "candidates": len(bids),
                "candidate_diagnostics": [
                    {
                        "title": b.title,
                        "url": b.source_url,
                        "domain": b.source_domain,
                        "scope": b.field("scope"),
                        "location": b.field("location"),
                        "due_date": b.field("due_date"),
                        "bid_number": b.field("bid_number"),
                        "contact_email": b.field("contact_email"),
                        "score": score_match(contractor, b).score,
                        "reasons": score_match(contractor, b).reasons,
                    }
                    for b in bids
                ],
            })
            continue
        safe = "".join(ch if ch.isalnum() or ch in "._-" else "_" for ch in contractor.company)[:60]
        folder = output / f"{index:04d}_{safe}"
        write_artifacts(matches[0], folder, matches[1:1 + args.related_count])
        manifest.append({"company": contractor.company, "status": "qualified", "score": matches[0].score, "bid": matches[0].bid.title, "source_url": matches[0].bid.source_url, "output": str(folder)})
    (output / "manifest.json").write_text(json.dumps({"contractor_audit": audit, "results": manifest}, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"contractors={len(contractors)} qualified={sum(x['status']=='qualified' for x in manifest)} output={output}")

def main() -> None:
    parser = argparse.ArgumentParser(prog="commercial-bid-pipeline")
    sub = parser.add_subparsers(dest="command", required=True)
    p = sub.add_parser("inspect-workbook")
    p.add_argument("--input", required=True)
    p.set_defaults(func=lambda a: inspect_workbook(a.input))
    p = sub.add_parser("run")
    p.add_argument("--input", required=True)
    p.add_argument("--output", default="out")
    p.add_argument("--state", default="")
    p.add_argument("--query", default="")
    p.add_argument("--limit", type=int, default=10)
    p.add_argument("--max-results", type=int, default=8)
    p.add_argument("--minimum-score", type=int, default=30)
    p.add_argument("--related-count", type=int, default=2)
    p.set_defaults(func=run)
    args = parser.parse_args()
    args.func(args)

if __name__ == "__main__":
    main()
